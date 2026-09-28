from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.time import local_naive_datetime
from app.models import Competition, Registration, RegistrationMember


EXPORT_HEADERS = (
    "Команда",
    "Капитан",
    "ФИО участника",
    "Почта",
    "Телефон",
    "Институт",
    "Курс",
    "Группа",
    "Дата заявки",
)
TITLE_ROW = 1
DETAILS_ROW = 2
SUMMARY_ROW = 3
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1
LAST_COLUMN = len(EXPORT_HEADERS)

TITLE_FILL = PatternFill(fill_type="solid", fgColor="084C8D")
DETAILS_FILL = PatternFill(fill_type="solid", fgColor="DCEEFF")
SUMMARY_FILL = PatternFill(fill_type="solid", fgColor="EEF5FB")
HEADER_FILL = PatternFill(fill_type="solid", fgColor="0D6EFD")
ALTERNATE_GROUP_FILL = PatternFill(fill_type="solid", fgColor="F3F8FD")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_BLUE = Side(style="thin", color="9CC7E8")
MEDIUM_BLUE = Side(style="medium", color="4B9BD3")


def build_competition_registrations_xlsx(competition: Competition) -> BytesIO:
    """Build a structured report with one row per registered participant."""
    workbook = Workbook()
    workbook.properties.title = f"Заявки — {competition.title}"
    workbook.properties.subject = "Заявки на соревнование"

    worksheet = workbook.active
    worksheet.title = "Заявки"
    _configure_worksheet(worksheet)
    _add_report_heading(worksheet, competition)
    _add_table_header(worksheet)

    registrations = sorted(
        competition.registrations,
        key=lambda item: (item.created_at, item.id),
    )
    for group_index, registration in enumerate(registrations):
        _add_registration_group(
            worksheet,
            registration,
            use_alternate_fill=group_index % 2 == 1,
        )

    _set_column_widths(worksheet)
    worksheet.print_area = f"A1:I{max(HEADER_ROW, worksheet.max_row)}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _configure_worksheet(worksheet: Worksheet) -> None:
    worksheet.freeze_panes = f"A{FIRST_DATA_ROW}"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = "0D6EFD"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    worksheet.oddFooter.center.text = "Страница &P из &N"
    worksheet.oddFooter.center.size = 9
    worksheet.oddFooter.center.color = "6C757D"


def _add_report_heading(
    worksheet: Worksheet,
    competition: Competition,
) -> None:
    for row in (TITLE_ROW, DETAILS_ROW, SUMMARY_ROW):
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=LAST_COLUMN,
        )

    title_cell = worksheet.cell(TITLE_ROW, 1, competition.title)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=18)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[TITLE_ROW].height = 34

    details_cell = worksheet.cell(DETAILS_ROW, 1, _competition_details(competition))
    details_cell.fill = DETAILS_FILL
    details_cell.font = Font(color="084C8D", bold=True, size=11)
    details_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    worksheet.row_dimensions[DETAILS_ROW].height = 28

    registration_count = len(competition.registrations)
    participant_count = sum(
        len(registration.members) for registration in competition.registrations
    )
    summary_cell = worksheet.cell(
        SUMMARY_ROW,
        1,
        f"Заявок: {registration_count}  •  Участников: {participant_count}",
    )
    summary_cell.fill = SUMMARY_FILL
    summary_cell.font = Font(color="425466", italic=True, size=10)
    summary_cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[SUMMARY_ROW].height = 22


def _competition_details(competition: Competition) -> str:
    registration_format = (
        "Командная регистрация"
        if competition.registration_type.value == "team"
        else "Индивидуальная регистрация"
    )
    dates = _date_range(competition.start_date, competition.end_date)
    return f"{registration_format}  •  {dates}  •  {competition.location}"


def _date_range(start: date, end: date) -> str:
    if start == end:
        return f"Дата проведения: {start:%d.%m.%Y}"
    return f"Даты проведения: {start:%d.%m.%Y}–{end:%d.%m.%Y}"


def _add_table_header(worksheet: Worksheet) -> None:
    for column, value in enumerate(EXPORT_HEADERS, start=1):
        cell = worksheet.cell(HEADER_ROW, column, value)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(
            left=THIN_BLUE,
            right=THIN_BLUE,
            top=THIN_BLUE,
            bottom=MEDIUM_BLUE,
        )
    worksheet.row_dimensions[HEADER_ROW].height = 30


def _add_registration_group(
    worksheet: Worksheet,
    registration: Registration,
    *,
    use_alternate_fill: bool,
) -> None:
    members = sorted(
        registration.members,
        key=lambda item: (not item.is_captain, item.id),
    )
    if not members:
        return

    first_row = worksheet.max_row + 1
    for member in members:
        worksheet.append(_member_row(registration, member))
        row = worksheet.max_row
        _style_participant_row(
            worksheet,
            row,
            is_first_member=row == first_row,
            is_captain=member.is_captain,
            use_alternate_fill=use_alternate_fill,
        )

    last_row = worksheet.max_row
    if last_row > first_row:
        for column in (1, 2, 9):
            worksheet.merge_cells(
                start_row=first_row,
                start_column=column,
                end_row=last_row,
                end_column=column,
            )

    for column in (1, 2, 9):
        worksheet.cell(first_row, column).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    worksheet.cell(first_row, 9).number_format = "DD.MM.YYYY HH:MM"


def _member_row(
    registration: Registration,
    member: RegistrationMember,
) -> tuple[str | int | datetime | None, ...]:
    return (
        registration.team_name or None,
        registration.captain_name,
        member.full_name,
        member.university_email,
        member.phone,
        member.institute,
        member.course,
        member.group_name,
        _excel_datetime(registration.created_at),
    )


def _style_participant_row(
    worksheet: Worksheet,
    row: int,
    *,
    is_first_member: bool,
    is_captain: bool,
    use_alternate_fill: bool,
) -> None:
    for column in range(1, LAST_COLUMN + 1):
        cell = worksheet.cell(row, column)
        if use_alternate_fill:
            cell.fill = ALTERNATE_GROUP_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(
            left=THIN_BLUE,
            right=THIN_BLUE,
            top=MEDIUM_BLUE if is_first_member else THIN_BLUE,
            bottom=THIN_BLUE,
        )
    if is_captain:
        for column in range(3, LAST_COLUMN):
            worksheet.cell(row, column).font = Font(bold=True)
    worksheet.row_dimensions[row].height = 30


def _excel_datetime(value: datetime) -> datetime:
    # Excel cannot store timezone-aware datetimes, so write Tyumen wall time.
    return local_naive_datetime(value)


def _set_column_widths(worksheet: Worksheet) -> None:
    widths = (24, 26, 28, 32, 20, 24, 10, 16, 20)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
