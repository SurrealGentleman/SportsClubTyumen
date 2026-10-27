from datetime import UTC, date, datetime, timedelta
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from tests.admin_helpers import admin_login

from app.core.security import get_password_hash
from app.models import (
    Admin,
    Competition,
    Registration,
    RegistrationMember,
    RegistrationType,
)
from app.services.excel_export import EXPORT_HEADERS


def create_admin(db_session: Session) -> None:
    db_session.add(
        Admin(
            login="admin",
            password_hash=get_password_hash("strong-admin-password"),
            is_active=True,
        )
    )
    db_session.commit()


def login(client: TestClient) -> None:
    response = admin_login(
        client,
        data={"username": "admin", "password": "strong-admin-password"},
    )
    assert response.status_code == 200


def create_competition(
    db_session: Session,
    *,
    registration_type: RegistrationType,
    slug: str,
) -> Competition:
    competition = Competition(
        title="Кубок ректора",
        slug=slug,
        location="ТюмГУ",
        start_date=date.today() + timedelta(days=20),
        end_date=date.today() + timedelta(days=20),
        registration_deadline=datetime.now(UTC) + timedelta(days=10),
        registration_type=registration_type,
        min_team_size=2 if registration_type == RegistrationType.team else None,
        max_team_size=5 if registration_type == RegistrationType.team else None,
        is_registration_open=True,
        is_published=True,
    )
    db_session.add(competition)
    db_session.flush()
    return competition


def test_admin_exports_one_row_per_team_member(
    client: TestClient,
    db_session: Session,
) -> None:
    create_admin(db_session)
    competition = create_competition(
        db_session,
        registration_type=RegistrationType.team,
        slug="rector-cup",
    )
    registration = Registration(
        competition=competition,
        team_name="Ладья",
        captain_name="Иван Капитанов",
        captain_email="captain@study.utmn.ru",
        captain_phone="+7 900 000-00-00",
        created_at=datetime(2026, 9, 23, 12, 30, tzinfo=UTC),
    )
    registration.members = [
        RegistrationMember(
            full_name="Пётр Участников",
            university_email="member@study.utmn.ru",
            institute="ИМиКН",
            course=2,
            group_name="02",
            is_captain=False,
        ),
        RegistrationMember(
            full_name="Иван Капитанов",
            university_email="captain@study.utmn.ru",
            institute="ИМиКН",
            course=3,
            group_name="01",
            phone="+7 900 000-00-00",
            is_captain=True,
        ),
    ]
    db_session.add(registration)
    db_session.commit()

    login(client)
    response = client.get(f"/admin/competition/action/export-xlsx?pks={competition.id}")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="rector-cup-registrations.xlsx"'
    )

    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook["Заявки"]
    header = tuple(
        worksheet.cell(row=5, column=column).value for column in range(1, 10)
    )

    assert worksheet["A1"].value == "Кубок ректора"
    assert "Командная регистрация" in worksheet["A2"].value
    assert "ТюмГУ" in worksheet["A2"].value
    assert worksheet["A3"].value == "Заявок: 1  •  Участников: 2"
    assert header == EXPORT_HEADERS
    assert worksheet.max_row == 7
    assert tuple(
        worksheet.cell(row=6, column=column).value for column in range(1, 9)
    ) == (
        "Ладья",
        "Иван Капитанов",
        "Иван Капитанов",
        "captain@study.utmn.ru",
        "+7 900 000-00-00",
        "ИМиКН",
        3,
        "01",
    )
    assert worksheet.cell(row=7, column=3).value == "Пётр Участников"
    assert worksheet.cell(row=7, column=5).value is None
    assert worksheet.cell(row=6, column=9).value == datetime(2026, 9, 23, 17, 30)
    assert {"A1:I1", "A2:I2", "A3:I3", "A6:A7", "B6:B7", "I6:I7"}.issubset(
        {str(cell_range) for cell_range in worksheet.merged_cells.ranges}
    )
    assert worksheet.freeze_panes == "A6"
    assert worksheet.page_setup.orientation == "landscape"


def test_individual_export_has_empty_team_column(
    client: TestClient,
    db_session: Session,
) -> None:
    create_admin(db_session)
    competition = create_competition(
        db_session,
        registration_type=RegistrationType.individual,
        slug="individual-cup",
    )
    registration = Registration(
        competition=competition,
        team_name=None,
        captain_name="Анна Спортсменова",
        captain_email="anna@study.utmn.ru",
        captain_phone="+7 900 111-22-33",
    )
    registration.members = [
        RegistrationMember(
            full_name="Анна Спортсменова",
            university_email="anna@study.utmn.ru",
            institute="ИГиП",
            course=1,
            group_name="11",
            phone="+7 900 111-22-33",
            is_captain=True,
        )
    ]
    db_session.add(registration)
    db_session.commit()

    login(client)
    response = client.get(f"/admin/competition/action/export-xlsx?pks={competition.id}")
    worksheet = load_workbook(BytesIO(response.content)).active

    assert response.status_code == 200
    assert worksheet.cell(row=1, column=1).value == "Кубок ректора"
    assert "Индивидуальная регистрация" in worksheet.cell(row=2, column=1).value
    assert worksheet.cell(row=6, column=1).value is None
    assert worksheet.cell(row=6, column=3).value == "Анна Спортсменова"


def test_excel_export_requires_admin_authentication(client: TestClient) -> None:
    response = client.get(
        "/admin/competition/action/export-xlsx?pks=1",
        follow_redirects=False,
    )

    assert response.status_code == 302
    assert response.headers["location"].endswith("/admin/login")
